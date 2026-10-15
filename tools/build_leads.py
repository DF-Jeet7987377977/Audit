"""Build leads.json for the DreamFolks Call Audit tool.

Universe  : Audit Tool/Talk Time Bands - Non-Workable, DNC & Rejected.xlsx -> Leads
Marketing : Sales Dashboard/Lead Data.xlsx -> Sheet1 (header on row 2)
Calls     : Sales Dashboard/Calling Data.xlsx -> Sheet1

Usage: python -I build_leads.py <out_dir>
Then upload <out_dir>/leads.json in the Call Audit tool: Settings > Lead pool.
"""
import io, json, re, subprocess, sys, base64, hashlib, datetime as dt
from collections import defaultdict
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
ROOT = r"C:/Users/SimranJeetSingh/OneDrive - DreamFolks Services Ltd/Desktop"
UNIVERSE = ROOT + "/Projects/Audit Tool/Talk Time Bands - Non-Workable, DNC & Rejected.xlsx"
LEADDATA = ROOT + "/Sales Dashboard/Lead Data.xlsx"
CALLING = ROOT + "/Sales Dashboard/Calling Data.xlsx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "."


def open_wb(path):
    """Open read-only; if Excel holds an exclusive lock, read through a shared .NET stream."""
    try:
        return openpyxl.load_workbook(path, read_only=True, data_only=True)
    except PermissionError:
        ps = ('$fs=[System.IO.File]::Open("%s","Open","Read","ReadWrite,Delete");'
              '$ms=New-Object System.IO.MemoryStream;$fs.CopyTo($ms);$fs.Close();'
              '[Convert]::ToBase64String($ms.ToArray())' % path.replace("/", "\\"))
        b64 = subprocess.run(["powershell.exe", "-NoProfile", "-Command", ps],
                             capture_output=True, text=True, check=True).stdout.strip()
        return openpyxl.load_workbook(io.BytesIO(base64.b64decode(b64)), read_only=True, data_only=True)


def phone10(v):
    d = re.sub(r"\D", "", str(v or ""))
    return d[-10:] if len(d) >= 10 else None


def lead_id(v):
    if v is None or v == "":
        return ""
    if isinstance(v, float):
        v = int(v)
    s = str(v).strip()
    s = s[:-2] if s.endswith(".0") else s
    return "" if s in ("0", "None", "nan") else s


def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str) and v.strip():
        for f in ("%d %b %Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
            try:
                return dt.datetime.strptime(v.strip(), f).date()
            except ValueError:
                pass
    return None


def clean(v):
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s in ("0", "None", "nan") else s


# ---------- universe ----------
ws = open_wb(UNIVERSE)["Leads"]
rows = ws.iter_rows(values_only=True)
hdr = [str(h).strip() if h else "" for h in next(rows)]
ix = {h: i for i, h in enumerate(hdr)}
universe = {}
dup = 0
for r in rows:
    p = phone10(r[ix["Phone Number"]])
    if not p:
        continue
    if p in universe:
        dup += 1
        continue
    universe[p] = {
        "lid": lead_id(r[ix["Lead ID"]]),
        "agent": clean(r[ix["Agent Name"]]) or "System",
        "created": as_date(r[ix["Lead Created"]]),
        "cat": clean(r[ix["Funnel Category"]]),
        "r1": clean(r[ix["Final Disposition (R1)"]]),
        "r2": clean(r[ix["Final Disposition (R2)"]]),
        "tt": round(float(r[ix["Talk Time (min)"]] or 0), 2),
        "band": clean(r[ix["Talk Time Category"]]),
        "dials": int(r[ix["Total Dials"]] or 0),
        "conn": clean(r[ix["Ever Connected"]]),
    }
print(f"universe leads: {len(universe)} (duplicate phones skipped: {dup})")
from collections import Counter
print("Lead ID formats:", Counter(("digits" if u["lid"].isdigit() else "other", len(u["lid"])) for u in universe.values() if u["lid"]).most_common(8))

# ---------- lead data (marketing) ----------
ws = open_wb(LEADDATA)["Sheet1"]
rows = ws.iter_rows(values_only=True)
next(rows)  # legends row
hdr = [str(h).strip() if h else "" for h in next(rows)]
ix = {h: i for i, h in enumerate(hdr) if h}
need = ["Phone Number", "Contact Source", "City", "Created On", "Lead ID", "Ad Name", "Ad set Name", "Campaign Name"]
missing = [n for n in need if n not in ix]
assert not missing, f"Lead Data missing columns: {missing}; have {list(ix)}"
by_lid, by_phone = {}, {}
for r in rows:
    p = phone10(r[ix["Phone Number"]])
    if not p or p not in universe:
        continue
    rec = {
        "src": clean(r[ix["Contact Source"]]),
        "city": clean(r[ix["City"]]).title(),
        "lcreated": as_date(r[ix["Created On"]]),
        "camp": clean(r[ix["Campaign Name"]]),
        "adset": clean(r[ix["Ad set Name"]]),
        "ad": clean(r[ix["Ad Name"]]),
    }
    lid = lead_id(r[ix["Lead ID"]])
    if lid:
        by_lid[lid] = rec
    prev = by_phone.get(p)
    if prev is None or (rec["lcreated"] or dt.date.min) >= (prev["lcreated"] or dt.date.min):
        by_phone[p] = rec

# ---------- calling data ----------
ws = open_wb(CALLING)["Sheet1"]
rows = ws.iter_rows(values_only=True)
hdr = [str(h).strip() if h else "" for h in next(rows)]
i_date, i_mob, i_status = hdr.index("Date Added"), hdr.index("Mobile Number"), hdr.index("Status")
calls = defaultdict(lambda: [0, 0, None, None])
for r in rows:
    p = phone10(r[i_mob])
    if not p or p not in universe:
        continue
    c = calls[p]
    c[0] += 1
    if str(r[i_status] or "").upper() == "CONNECTED":
        c[1] += 1
    d = as_date(r[i_date])
    if d:
        c[2] = d if c[2] is None or d < c[2] else c[2]
        c[3] = d if c[3] is None or d > c[3] else c[3]

# ---------- assemble (dictionary-encoded) ----------
dicts = {k: [] for k in ("agent", "cat", "r1", "r2", "band", "src", "city", "camp", "adset", "ad")}
lookup = {k: {} for k in dicts}


def enc(k, v):
    m = lookup[k]
    if v not in m:
        m[v] = len(dicts[k])
        dicts[k].append(v)
    return m[v]


# Masking: phone numbers never appear in the published file.
# ph = PBKDF2-SHA256(phone10, SALT, ITER, 12 bytes); lh = same over the Lead ID;
# le = Lead ID XOR PBKDF2-SHA256(phone10, SALT+":lid", ITER, len(lid)), so the page recovers it once the auditor types the phone.
SALT, ITER = b"dreamfolks-call-audit-v1", 10000


def mask(v):
    return hashlib.pbkdf2_hmac("sha256", v.encode(), SALT, ITER, dklen=12).hex()


def seal_lid(lid, phone):
    if not lid:
        return ""
    data = lid.encode()
    key = hashlib.pbkdf2_hmac("sha256", phone.encode(), SALT + b":lid", ITER, dklen=len(data))
    return bytes(a ^ b for a, b in zip(data, key)).hex()


def iso(d):
    return d.isoformat() if d else ""


FIELDS = ["ph", "lh", "le", "agent", "created", "cat", "r1", "r2", "tt", "band", "dials", "conn",
          "src", "city", "camp", "adset", "ad", "calls", "ccalls", "first", "last"]
out_rows = []
hit_lid = hit_phone = hit_calls = 0
for p, u in universe.items():
    m = by_lid.get(u["lid"]) if u["lid"] else None
    if m:
        hit_lid += 1
    else:
        m = by_phone.get(p)
        if m:
            hit_phone += 1
    m = m or {"src": "", "city": "", "lcreated": None, "camp": "", "adset": "", "ad": ""}
    c = calls.get(p)
    if c:
        hit_calls += 1
    c = c or [0, 0, None, None]
    out_rows.append([
        mask(p), mask(u["lid"]) if u["lid"] else "", seal_lid(u["lid"], p), enc("agent", u["agent"]), iso(u["created"] or m["lcreated"]),
        enc("cat", u["cat"]), enc("r1", u["r1"]), enc("r2", u["r2"]), u["tt"], enc("band", u["band"]),
        u["dials"], 1 if u["conn"].lower().startswith("y") else 0,
        enc("src", m["src"]), enc("city", m["city"]), enc("camp", m["camp"]),
        enc("adset", m["adset"]), enc("ad", m["ad"]),
        c[0], c[1], iso(c[2]), iso(c[3]),
    ])

n = len(out_rows)
print(f"Lead Data match: {hit_lid} by Lead ID + {hit_phone} by phone = {hit_lid + hit_phone}/{n} "
      f"({(hit_lid + hit_phone) / n:.1%})")
print(f"Calling Data match: {hit_calls}/{n} ({hit_calls / n:.1%})")
assert len({r[0] for r in out_rows}) == n, "phone hash collision"
payload = {"built": dt.datetime.now().strftime("%d %b %Y %H:%M"), "mask": {"salt": SALT.decode(), "iter": ITER, "bytes": 12}, "fields": FIELDS, "dicts": dicts, "rows": out_rows}
with open(OUT + "/leads.json", "w", encoding="utf-8") as f:
    json.dump(payload, f, ensure_ascii=False, separators=(",", ":"))
print("dict sizes:", {k: len(v) for k, v in dicts.items()})
for row in out_rows[:1]:
    print({FIELDS[i]: (dicts[FIELDS[i]][v] if FIELDS[i] in dicts else v) for i, v in enumerate(row)})
