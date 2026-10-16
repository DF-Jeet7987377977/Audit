"""Build the masked lead pool from the source Excel files.

  leads     (required) "Talk Time Bands - Non-Workable, DNC & Rejected.xlsx", sheet Leads: the leads to audit
  leaddata  (optional) "Lead Data.xlsx": source, city, campaign, ad set, ad, created date
  calling   (optional) "Calling Data.xlsx": total / connected calls, first and last call date

Phone numbers and Lead IDs are masked the same way the page expects:
  ph = PBKDF2-SHA256(phone10, SALT, ITER, 12 bytes)        lh = same over the Lead ID
  le = Lead ID XOR PBKDF2-SHA256(phone10, SALT+":lid", ITER, len(Lead ID))
"""
import datetime as dt
import hashlib
import re
from collections import defaultdict

import openpyxl

SALT, ITER = b"dreamfolks-call-audit-v1", 10000
FIELDS = ["ph", "lh", "le", "agent", "created", "cat", "r1", "r2", "tt", "band", "dials", "conn",
          "src", "city", "camp", "adset", "ad", "calls", "ccalls", "first", "last"]
LEADS_COLS = ["Phone Number", "Lead ID", "Agent Name", "Lead Created", "Funnel Category", "Final Disposition (R1)",
              "Final Disposition (R2)", "Talk Time (min)", "Talk Time Category", "Total Dials", "Ever Connected"]
LEADDATA_COLS = ["Phone Number", "Contact Source", "City", "Created On", "Lead ID", "Ad Name", "Ad set Name", "Campaign Name"]
CALLING_COLS = ["Date Added", "Mobile Number", "Status"]


def phone10(v):
    d = re.sub(r"\D", "", str(v or ""))
    return d[-10:] if len(d) >= 10 else None


def lead_id(v):
    if v is None or v == "":
        return ""
    if isinstance(v, float):
        v = int(v)
    s = str(v).strip()
    s = s[:-2] if s.endswith(".0") else s
    return "" if s in ("0", "None", "nan") else s


def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str) and v.strip():
        for f in ("%d %b %Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y"):
            try:
                return dt.datetime.strptime(v.strip(), f).date()
            except ValueError:
                pass
    return None


def clean(v):
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s in ("0", "None", "nan") else s


def num(v, cast=float):
    try:
        return cast(v or 0)
    except (TypeError, ValueError):
        return cast(0)


def find_table(path, need, label, prefer=None):
    """Open a workbook and find the sheet + header row (within the first 10 rows) holding all `need` columns."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = ([prefer] if prefer in wb.sheetnames else []) + [n for n in wb.sheetnames if n != prefer]
    for name in sheets:
        rows = wb[name].iter_rows(values_only=True)
        for _ in range(10):
            try:
                hdr = [str(h).strip() if h is not None else "" for h in next(rows)]
            except StopIteration:
                break
            ix = {}
            for i, h in enumerate(hdr):
                ix.setdefault(h, i)          # first column wins when a header repeats (Calling Data has two "Status")
            if all(c in ix for c in need):
                return ix, rows
    raise ValueError(f"{label}: no sheet has all of these columns: {', '.join(need)}")


def build(leads_path, leaddata_path=None, calling_path=None, progress=lambda msg: None):
    progress("Reading the audit leads file…")
    ix, rows = find_table(leads_path, LEADS_COLS, "Audit leads file", prefer="Leads")
    universe, dup = {}, 0
    for r in rows:
        p = phone10(r[ix["Phone Number"]])
        if not p:
            continue
        if p in universe:
            dup += 1
            continue
        universe[p] = {
            "lid": lead_id(r[ix["Lead ID"]]), "agent": clean(r[ix["Agent Name"]]) or "System",
            "created": as_date(r[ix["Lead Created"]]), "cat": clean(r[ix["Funnel Category"]]),
            "r1": clean(r[ix["Final Disposition (R1)"]]), "r2": clean(r[ix["Final Disposition (R2)"]]),
            "tt": round(num(r[ix["Talk Time (min)"]]), 2), "band": clean(r[ix["Talk Time Category"]]),
            "dials": num(r[ix["Total Dials"]], int), "conn": clean(r[ix["Ever Connected"]]),
        }
    if not universe:
        raise ValueError("Audit leads file: no rows with a phone number were found.")

    by_lid, by_phone = {}, {}
    if leaddata_path:
        progress(f"Matching {len(universe):,} leads to Lead Data…")
        ix, rows = find_table(leaddata_path, LEADDATA_COLS, "Lead Data file", prefer="Sheet1")
        for r in rows:
            p = phone10(r[ix["Phone Number"]])
            if not p or p not in universe:
                continue
            rec = {"src": clean(r[ix["Contact Source"]]), "city": clean(r[ix["City"]]).title(),
                   "lcreated": as_date(r[ix["Created On"]]), "camp": clean(r[ix["Campaign Name"]]),
                   "adset": clean(r[ix["Ad set Name"]]), "ad": clean(r[ix["Ad Name"]])}
            lid = lead_id(r[ix["Lead ID"]])
            if lid:
                by_lid[lid] = rec
            prev = by_phone.get(p)
            if prev is None or (rec["lcreated"] or dt.date.min) >= (prev["lcreated"] or dt.date.min):
                by_phone[p] = rec

    calls = defaultdict(lambda: [0, 0, None, None])
    if calling_path:
        progress("Counting calls in Calling Data (large file, this takes a few minutes)…")
        ix, rows = find_table(calling_path, CALLING_COLS, "Calling Data file", prefer="Sheet1")
        i_date, i_mob, i_status = ix["Date Added"], ix["Mobile Number"], ix["Status"]
        for r in rows:
            p = phone10(r[i_mob])
            if not p or p not in universe:
                continue
            c = calls[p]
            c[0] += 1
            if str(r[i_status] or "").upper() == "CONNECTED":
                c[1] += 1
            d = as_date(r[i_date])
            if d:
                c[2] = d if c[2] is None or d < c[2] else c[2]
                c[3] = d if c[3] is None or d > c[3] else c[3]

    dicts = {k: [] for k in ("agent", "cat", "r1", "r2", "band", "src", "city", "camp", "adset", "ad")}
    lookup = {k: {} for k in dicts}

    def enc(k, v):
        m = lookup[k]
        if v not in m:
            m[v] = len(dicts[k])
            dicts[k].append(v)
        return m[v]

    def mask(v):
        return hashlib.pbkdf2_hmac("sha256", v.encode(), SALT, ITER, dklen=12).hex()

    def seal(lid, phone):
        if not lid:
            return ""
        data = lid.encode()
        key = hashlib.pbkdf2_hmac("sha256", phone.encode(), SALT + b":lid", ITER, dklen=len(data))
        return bytes(a ^ b for a, b in zip(data, key)).hex()

    iso = lambda d: d.isoformat() if d else ""
    out, hit_lead, hit_calls = [], 0, 0
    empty = {"src": "", "city": "", "lcreated": None, "camp": "", "adset": "", "ad": ""}
    for n, (p, u) in enumerate(universe.items()):
        if n % 1000 == 0:
            progress(f"Masking phone numbers and Lead IDs… {n:,} of {len(universe):,}")
        m = (by_lid.get(u["lid"]) if u["lid"] else None) or by_phone.get(p)
        hit_lead += bool(m)
        m = m or empty
        c = calls.get(p)
        hit_calls += bool(c)
        c = c or [0, 0, None, None]
        out.append([
            mask(p), mask(u["lid"]) if u["lid"] else "", seal(u["lid"], p), enc("agent", u["agent"]), iso(u["created"] or m["lcreated"]),
            enc("cat", u["cat"]), enc("r1", u["r1"]), enc("r2", u["r2"]), u["tt"], enc("band", u["band"]),
            u["dials"], 1 if u["conn"].lower().startswith("y") else 0,
            enc("src", m["src"]), enc("city", m["city"]), enc("camp", m["camp"]), enc("adset", m["adset"]), enc("ad", m["ad"]),
            c[0], c[1], iso(c[2]), iso(c[3]),
        ])
    payload = {"built": dt.datetime.now().strftime("%d %b %Y %H:%M"), "mask": {"salt": SALT.decode(), "iter": ITER, "bytes": 12},
               "fields": FIELDS, "dicts": dicts, "rows": out}
    stats = {"leads": len(out), "duplicatePhones": dup,
             "leadDataMatched": hit_lead if leaddata_path else None, "callingMatched": hit_calls if calling_path else None}
    return payload, stats
